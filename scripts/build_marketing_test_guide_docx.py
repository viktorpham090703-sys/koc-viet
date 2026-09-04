# -*- coding: utf-8 -*-
"""
Script to generate the Marketing-friendly User Testing Guide for KOC Viet in .docx format.
Target Audience: Marketing team / Non-technical Testers.
Focuses strictly on end-to-end user journeys, screen navigation, and verification steps.
"""

from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"D:\koc-viet")
OUTPUT_DOCX = ROOT / "docs" / "Huong-dan-kiem-thu-KOC-Viet-production.docx"

# Color Palette
NAVY = "17365D"          # Deep Navy for primary headings, table headers
BLUE = "2563EB"          # Vibrant Brand Blue for subheadings and action cues
LIGHT_BLUE = "F0F7FD"    # Light blue for table alternate rows & callout
DARK_GRAY = "1F2937"     # Dark slate for body text (high contrast)
MID_GRAY = "64748B"      # Muted gray for captions & header/footer
BORDER_GRAY = "CBD5E1"   # Subtle borders
SUCCESS_GREEN = "059669" # Green for tips & pass notes
SUCCESS_BG = "F0FDF4"
WARN_AMBER = "D97706"    # Amber for important notes & warnings
WARN_BG = "FFFBEB"
WHITE = "FFFFFF"

def set_run_font(run, name="Calibri", size=10.5, color=DARK_GRAY, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_margins(cell, top=100, bottom=100, start=140, end=140):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")

def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)

def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths[min(idx, len(widths) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)

def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Trang ")
    set_run_font(run, size=9, color=MID_GRAY)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])

def configure(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.27)   # A4 Width
    section.page_height = Inches(11.69) # A4 Height
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.4)
    section.footer_distance = Inches(0.4)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(DARK_GRAY)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, color, before, after in (
        ("Heading 1", 15, NAVY, 14, 5),
        ("Heading 2", 12.5, BLUE, 10, 4),
        ("Heading 3", 11, NAVY, 7, 3),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    # Header & Footer
    header_p = section.header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run_font(header_p.add_run("KOC VIỆT  |  HƯỚNG DẪN KIỂM THỬ TRẢI NGHIỆM NGƯỜI DÙNG (USER TEST)"), size=8.5, color=MID_GRAY, bold=True)

    footer_p = section.footer.paragraphs[0]
    set_run_font(footer_p.add_run("Tài liệu nội bộ KOC Việt • https://kocviet.com/  |  "), size=9, color=MID_GRAY)
    add_page_number(footer_p)

def add_callout(doc, title, text, fill=LIGHT_BLUE, border_color=BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    width_dxa = 9360
    
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=120, bottom=120, start=160, end=160)
    
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = OxmlElement("w:tcBorders")
    left_bdr = OxmlElement("w:left")
    left_bdr.set(qn("w:val"), "single")
    left_bdr.set(qn("w:sz"), "24") # 3pt
    left_bdr.set(qn("w:space"), "0")
    left_bdr.set(qn("w:color"), border_color)
    tc_borders.append(left_bdr)
    
    for side in ("top", "bottom", "right"):
        bdr = OxmlElement(f"w:{side}")
        bdr.set(qn("w:val"), "none")
        tc_borders.append(bdr)
    tc_pr.append(tc_borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    set_run_font(p.add_run(f"{title}\n"), size=10.5, color=NAVY, bold=True)
    set_run_font(p.add_run(text), size=10, color=DARK_GRAY)
    
    set_table_geometry(table, [width_dxa])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def add_heading(doc, text, level=1):
    doc.add_paragraph(text, style=f"Heading {level}")

def add_para(doc, text, bold_prefix=None, italic=False):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        set_run_font(p.add_run(bold_prefix), size=10.5, color=DARK_GRAY, bold=True)
        set_run_font(p.add_run(text[len(bold_prefix):]), size=10.5, color=DARK_GRAY, italic=italic)
    else:
        set_run_font(p.add_run(text), size=10.5, color=DARK_GRAY, italic=italic)
    return p

def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2.5)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix and text.startswith(bold_prefix):
        set_run_font(p.add_run(bold_prefix), size=10.5, color=DARK_GRAY, bold=True)
        set_run_font(p.add_run(text[len(bold_prefix):]), size=10.5, color=DARK_GRAY)
    else:
        set_run_font(p.add_run(text), size=10.5, color=DARK_GRAY)
    return p

def add_step(doc, number_str, title, details):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    set_run_font(p.add_run(f"▶ {number_str}: {title}\n"), size=10.5, color=BLUE, bold=True)
    set_run_font(p.add_run(details), size=10, color=DARK_GRAY)
    return p

def add_table(doc, headers, rows, widths, center_cols=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    if center_cols is None:
        center_cols = [0]
    
    # Header row
    for idx, header in enumerate(headers):
        cell = table.cell(0, idx)
        cell.text = ""
        set_cell_shading(cell, NAVY)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(p.add_run(header), size=10, color=WHITE, bold=True)
    set_repeat_table_header(table.rows[0])
    
    # Data rows
    for r_idx, row in enumerate(rows, start=1):
        is_even = (r_idx % 2 == 0)
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = ""
            if is_even:
                set_cell_shading(cell, LIGHT_BLUE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            
            if c_idx in center_cols or len(str(val)) <= 8:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                
            set_run_font(p.add_run(str(val)), size=9.5, color=DARK_GRAY)
            set_cell_margins(cell, top=70, bottom=70, start=110, end=110)
            
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def build_marketing_guide():
    doc = Document()
    configure(doc)

    # ----------------------------------------------------
    # COVER / TITLE BLOCK
    # ----------------------------------------------------
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(12)
    title_p.paragraph_format.space_after = Pt(4)
    title_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run_font(title_p.add_run("HƯỚNG DẪN KIỂM THỬ TRẢI NGHIỆM NGƯỜI DÙNG (USER TEST)"), size=20, color=NAVY, bold=True)

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(10)
    set_run_font(sub_p.add_run("Cẩm nang kiểm tra các luồng nghiệp vụ chính dành cho đội ngũ Marketing & Vận hành sàn KOC Việt"), size=12, color=BLUE, italic=True)

    add_callout(
        doc,
        "TÓM TẮT DÀNH CHO BẠN TESTER MARKETING",
        "• Đối tượng áp dụng: Các bạn chuyên viên Marketing, Quản lý chiến dịch và Tester trải nghiệm nghiệp vụ.\n"
        "• Hệ thống kiểm thử: Website nền tảng KOC Việt (Địa chỉ: https://kocviet.com/).\n"
        "• Tinh thần tài liệu: Hướng dẫn thực chiến 100% thao tác màn hình (Bấm vào đâu — Điền thông tin gì — Màn hình hiện kết quả ra sao), loại bỏ hoàn toàn các câu lệnh lập trình, mã lỗi kỹ thuật hoặc truy vấn dữ liệu phức tạp.\n"
        "• Mục tiêu: Giúp bạn tự tin nhập vai Doanh nghiệp (Brand), KOC và Quản trị viên (Admin) để kiểm tra toàn bộ luồng vận hành thực tế.",
        fill=LIGHT_BLUE,
        border_color=BLUE
    )

    # ----------------------------------------------------
    # PHẦN 1: CHUẨN BỊ TRƯỚC KHI BẮT ĐẦU KIỂM THỬ
    # ----------------------------------------------------
    add_heading(doc, "1. Chuẩn bị trước khi bắt đầu kiểm thử", 1)
    
    add_para(doc, "Để quá trình kiểm thử diễn ra liền mạch, bạn chỉ cần chuẩn bị một số công cụ và thông tin cơ bản sau:")
    
    add_bullet(doc, "Địa chỉ website cần mở: https://kocviet.com/ (truy cập bằng trình duyệt Chrome, Edge hoặc Cốc Cốc trên máy tính).", "• Địa chỉ website cần mở: ")
    add_bullet(doc, "Thiết bị đề xuất: Máy tính (Laptop/PC) để thao tác dễ dàng trên giao diện Doanh nghiệp và Admin. Bạn cũng có thể dùng thêm điện thoại di động để trải nghiệm giao diện KOC chuẩn Mobile.", "• Thiết bị đề xuất: ")
    add_bullet(doc, "Mẹo mở song song nhiều vai trò: Khi test các tính năng phối hợp 2 chiều (ví dụ: Doanh nghiệp gửi booking ➔ KOC nhận việc), bạn nên mở 2 hoặc 3 trình duyệt khác nhau:", "• Mẹo mở song song nhiều vai trò: ")
    
    # Bảng phân chia trình duyệt
    add_table(
        doc,
        ["Trình duyệt / Cửa sổ", "Vai trò đăng nhập", "Mục đích sử dụng chính khi test"],
        [
            ("Google Chrome bình thường", "Doanh nghiệp (Brand)", "Tìm kiếm KOC, tạo booking, duyệt video, đánh giá và quản lý chiến dịch."),
            ("Microsoft Edge hoặc Cốc Cốc", "KOC (Creator)", "Đăng ký kênh, nhận đơn booking, nộp video demo, dán link bài đăng, rút tiền."),
            ("Cửa sổ ẩn danh (Incognito)", "Quản trị viên (Admin)", "Duyệt tài khoản KOC, duyệt doanh nghiệp, gửi báo giá, giải ngân và xử lý khiếu nại.")
        ],
        [2800, 2400, 4160],
        center_cols=[0, 1]
    )

    add_callout(
        doc,
        "LƯU Ý VỀ DỮ LIỆU TEST TRÊN PRODUCTION",
        "1. Đặt tên dữ liệu mẫu: Khi tạo KOC, sản phẩm, chiến dịch hoặc booking thử nghiệm, hãy đặt tên bắt đầu bằng tiền tố [TEST-MKT] (Ví dụ: [TEST-MKT] Review Kem Chống Nắng) để cả nhóm dễ dàng phân biệt và dọn dẹp sau khi kiểm thử.\n"
        "2. Không dùng tiền thật: Không chuyển tiền thật từ tài khoản ngân hàng cá nhân hoặc thanh toán số tiền lớn ngoài kế hoạch kiểm thử đã được công ty phê duyệt.",
        fill=WARN_BG,
        border_color=WARN_AMBER
    )

    # ----------------------------------------------------
    # PHẦN 2: LUỒNG 1 — KOC ĐĂNG KÝ & KÍCH HOẠT HỒ SƠ
    # ----------------------------------------------------
    add_heading(doc, "2. Luồng 1 — KOC Đăng ký & Kích hoạt hồ sơ (Creator Onboarding)", 1)
    
    add_para(doc, "Mục tiêu: Đóng vai một bạn sáng tạo nội dung mới tham gia sàn, trải nghiệm toàn bộ hành trình tạo tài khoản, xác minh kênh mạng xã hội, ký hợp đồng và được Admin phê duyệt.")

    add_step(doc, "Bước 1", "Mở trang đăng ký & Xác thực Email", 
             "• Truy cập đường dẫn: https://kocviet.com/#/tuyen-koc (hoặc bấm nút 'Trở thành KOC' ở góc trên trang chủ).\n"
             "• Nhập địa chỉ Email cá nhân chưa từng đăng ký trên hệ thống ➔ Bấm nút 'Nhận mã xác nhận'.\n"
             "• Mở hòm thư email của bạn, lấy mã OTP gồm 6 chữ số và nhập vào ô xác thực trên màn hình.")

    add_step(doc, "Bước 2", "Điền thông tin cá nhân & Kênh mạng xã hội",
             "• Nhập họ tên, số điện thoại, tỉnh/thành phố nơi bạn đang sinh sống.\n"
             "• Chọn nền tảng mạng xã hội chính của bạn: TikTok, Facebook, Instagram, YouTube hoặc Threads.\n"
             "• Nhập số lượng người theo dõi (Follower) thực tế trên kênh; tài khoản KOC cần có tối thiểu 1.000 người theo dõi.")

    add_step(doc, "Bước 3", "Kiểm tra thông tin kênh và số người theo dõi",
             "• Kiểm tra link kênh mạng xã hội đã nhập đúng và có thể truy cập.\n"
             "• Kiểm tra tổng số người theo dõi là số nguyên từ 1.000 trở lên và đúng với thông tin hiện tại trên kênh.\n"
             "➔ Kết quả mong đợi: Hệ thống ghi nhận thông tin và chuyển hồ sơ sang bước phân hạng; Admin sẽ kiểm duyệt trước khi kích hoạt.")

    add_step(doc, "Bước 4", "Thiết lập bảng giá dịch vụ theo hạng KOC",
             "• Dựa vào số người theo dõi, hệ thống tự động xếp hạng bạn vào nhóm phù hợp (Nano, Micro, Mid, Macro hoặc Mega).\n"
             "• Điền mức giá mong muốn cho từng gói dịch vụ (Video Review, Video Quảng cáo). Hệ thống có hiển thị khoảng giá gợi ý để bạn tham khảo.")

    add_step(doc, "Bước 5", "Điền thông tin nhận tiền & Ký hợp đồng online",
             "• Chọn tên ngân hàng, nhập số tài khoản ngân hàng và tên chủ tài khoản.\n"
             "• Đọc nhanh các điều khoản hợp tác trên màn hình.\n"
             "• Ký tên điện tử trực tiếp bằng chuột hoặc chạm tay trên màn hình điện thoại ➔ Bấm 'Hoàn tất đăng ký'.\n"
             "➔ Màn hình thông báo: Hồ sơ của bạn đã được gửi thành công và đang chờ Ban Quản trị phê duyệt.")

    add_step(doc, "Bước 6", "Admin phê duyệt hồ sơ & Nghiệm thu hiển thị",
             "• Chuyển sang trình duyệt đăng nhập Admin ➔ Vào menu 'Quản lý KOC' (#/queue).\n"
             "• Tìm tên bạn KOC vừa đăng ký ➔ Bấm xem chi tiết hồ sơ và bấm nút 'Duyệt hồ sơ'.\n"
             "• Kiểm tra kết quả: Đăng nhập tài khoản KOC thấy vào thẳng màn hình làm việc (#/home). Mở tài khoản Doanh nghiệp vào mục 'Tìm KOC' (#/find) sẽ thấy hồ sơ bạn KOC mới xuất hiện đẹp mắt trên sàn.")

    add_callout(
        doc,
        "MẸO KIỂM TRA BẮT LỖI (DÀNH CHO TESTER)",
        "Hãy thử nhập link kênh không hợp lệ hoặc số follower dưới 1.000. Hệ thống phải hiển thị thông báo lỗi bằng tiếng Việt rõ ràng và không cho phép sang bước kế tiếp.",
        fill=SUCCESS_BG,
        border_color=SUCCESS_GREEN
    )

    # ----------------------------------------------------
    # PHẦN 3: LUỒNG 2 — DOANH NGHIỆP ĐĂNG KÝ TÀI KHOẢN
    # ----------------------------------------------------
    add_heading(doc, "3. Luồng 2 — Doanh nghiệp Đăng ký tài khoản (Brand Onboarding)", 1)
    
    add_para(doc, "Mục tiêu: Đóng vai một Nhãn hàng / Doanh nghiệp mới đăng ký tài khoản để bắt đầu tuyển chọn KOC cho chiến dịch marketing của công ty.")

    add_step(doc, "Bước 1", "Nhập thông tin đăng ký Doanh nghiệp",
             "• Truy cập đường dẫn: https://kocviet.com/#/business-register (hoặc bấm 'Đăng ký doanh nghiệp' tại trang đăng nhập).\n"
             "• Nhập địa chỉ Email công ty ➔ Nhận mã OTP qua email và nhập xác thực.\n"
             "• Điền thông tin công ty: Tên công ty / Tên nhãn hàng, Người đại diện liên hệ, Số điện thoại, Mã số thuế doanh nghiệp.\n"
             "• Tải lên ảnh Giấy phép đăng ký kinh doanh (dùng file ảnh mẫu để kiểm thử) ➔ Bấm nút 'Gửi đăng ký'.\n"
             "➔ Kết quả mong đợi: Hệ thống hiển thị thông báo gửi hồ sơ thành công và hồ sơ chuyển sang trạng thái chờ duyệt.")

    add_step(doc, "Bước 2", "Admin phê duyệt tài khoản Doanh nghiệp",
             "• Đăng nhập tài khoản Admin ➔ Mở menu 'Quản lý doanh nghiệp' (#/businesses).\n"
             "• Tìm tên doanh nghiệp vừa tạo trong danh sách 'Chờ duyệt' ➔ Bấm nút 'Phê duyệt'.")

    add_step(doc, "Bước 3", "Nghiệm thu đăng nhập Doanh nghiệp",
             "• Quay lại trình duyệt Doanh nghiệp, đăng nhập bằng Email và Mật khẩu vừa tạo.\n"
             "➔ Kết quả mong đợi: Đăng nhập thành công và dẫn vào thẳng Trang tổng quan quản lý (#/dashboard) với đầy đủ thanh menu bên trái: Tìm KOC, Booking, Ví, Chiến dịch lớn...")

    # ----------------------------------------------------
    # PHẦN 4: LUỒNG 3 — QUY TRÌNH BOOKING KOC 1-1 TRÊN CHỢ MARKETPLACE
    # ----------------------------------------------------
    add_heading(doc, "4. Luồng 3 — Booking KOC 1-1 trên Chợ Marketplace (Luồng cốt lõi)", 1)
    
    add_para(doc, "Đây là tính năng cốt lõi và quan trọng nhất của KOC Việt. Hãy thực hiện từng bước phối hợp nhịp nhàng giữa Doanh nghiệp và KOC theo bảng tóm tắt dưới đây:")

    # Bảng 7 bước Booking
    add_table(
        doc,
        ["Bước", "Bên thao tác", "Màn hình", "Thao tác chi tiết", "Kết quả hiển thị trên màn hình"],
        [
            ("1", "Doanh nghiệp", "#/find (Tìm KOC)", "Tìm bạn KOC mong muốn, bấm vào xem hồ sơ và bấm nút 'Đặt lịch / Booking'.", "Mở ra biểu mẫu tạo yêu cầu booking với KOC."),
            ("2", "Doanh nghiệp", "#/find (Đặt lịch)", "Chọn gói Review/Quảng cáo, dán link sản phẩm, nhập kịch bản (brief) và hạn chót nộp video (deadline) ➔ Bấm 'Xác nhận'.", "Đơn tạo thành công; Số tiền booking được tạm giữ trong quỹ Escrow (chưa chuyển cho KOC ngay)."),
            ("3", "KOC", "#/bookings (Booking)", "Mở danh sách đơn hàng, thấy đơn mới ở trạng thái 'Chờ xác nhận' ➔ Bấm nút 'Nhận việc'.", "Trạng thái đơn hàng chuyển sang 'Đang thực hiện / Chờ nộp video'."),
            ("4", "KOC", "#/bookings (Chi tiết)", "Bấm 'Bắt đầu làm' ➔ Tải file video demo/nháp lên (hoặc dán link video xem trước) ➔ Bấm 'Gửi duyệt video'.", "Trạng thái đơn hàng chuyển sang 'Chờ Doanh nghiệp duyệt video'."),
            ("5A", "Doanh nghiệp", "#/orders (Đơn hàng)", "(Nếu video chưa đạt) Bấm nút 'Yêu cầu sửa', điền góp ý chi tiết ➔ KOC tải lên bản dựng mới.", "KOC nhận được góp ý và có nút tải lại video sửa đổi."),
            ("5B", "Doanh nghiệp", "#/orders (Đơn hàng)", "(Khi video đã đạt) Bấm nút 'Duyệt video'.", "Trạng thái đơn chuyển sang 'Video đã duyệt, chờ KOC đăng bài'."),
            ("6", "KOC", "#/bookings (Chi tiết)", "Đăng video lên kênh MXH thật ➔ Dán link bài viết chính thức vào hệ thống ➔ Bấm 'Gửi bài đăng'.", "Trạng thái đơn chuyển sang 'Đã đăng bài, chờ Brand nghiệm thu'."),
            ("7", "Doanh nghiệp", "#/orders (Chi tiết)", "Bấm mở link bài đăng xem thử ➔ Bấm nút 'Hoàn tất đơn hàng' ➔ Chọn số sao (1-5 sao) và viết nhận xét.", "Đơn hàng hoàn tất; Tiền thù lao được giải ngân vào ví KOC; Đánh giá sao hiển thị trên hồ sơ KOC.")
        ],
        [600, 1500, 1600, 2960, 2700],
        center_cols=[0, 1, 2]
    )

    add_callout(
        doc,
        "ĐIỂM KIỂM TRA ĐẶC BIỆT VỀ DÒNG TIỀN (ESCROW)",
        "• Khi Doanh nghiệp bấm tạo đơn: Tiền trong ví của Doanh nghiệp bị trừ khỏi 'Số dư khả dụng' và chuyển sang 'Đang tạm giữ' (Escrow). Điều này giúp KOC yên tâm sản xuất mà không sợ bùng tiền.\n"
        "• Nếu KOC bấm 'Từ chối nhận việc': Toàn bộ số tiền tạm giữ sẽ tự động hoàn trả 100% về số dư khả dụng của Doanh nghiệp ngay lập tức.\n"
        "• Chỉ khi Doanh nghiệp bấm 'Hoàn tất đơn hàng': Tiền tạm giữ mới được giải phóng để chuyển vào ví KOC (sau khi sàn trích lại phí dịch vụ minh bạch).",
        fill=LIGHT_BLUE,
        border_color=BLUE
    )

    # ----------------------------------------------------
    # PHẦN 5: LUỒNG 4 — QUẢN LÝ VÍ, NẠP TIỀN & RÚT TIỀN
    # ----------------------------------------------------
    add_heading(doc, "5. Luồng 4 — Quản lý Ví, Nạp tiền & Rút tiền", 1)
    
    add_para(doc, "Hệ thống ví điện tử nội bộ giúp Doanh nghiệp chủ động ngân sách và giúp KOC nhận thù lao nhanh chóng.")

    add_heading(doc, "5.1. Doanh nghiệp nạp tiền vào ví (Qua cổng PayOS)", 2)
    add_bullet(doc, "Vào menu 'Ví doanh nghiệp' (#/wallet) trên tài khoản Doanh nghiệp.", "• Bước 1: ")
    add_bullet(doc, "Xem số dư khả dụng hiện tại ➔ Bấm nút 'Nạp tiền'.", "• Bước 2: ")
    add_bullet(doc, "Nhập số tiền muốn nạp thử nghiệm (ví dụ: 100.000 VNĐ) ➔ Bấm 'Tạo yêu cầu nạp'.", "• Bước 3: ")
    add_bullet(doc, "Màn hình xuất hiện mã QR PayOS thanh toán. Bạn quét thử nghiệm hoặc mở thanh toán test.", "• Bước 4: ")
    add_bullet(doc, "Kiểm tra kết quả: Màn hình cập nhật trạng thái 'Thanh toán thành công' và số dư ví khả dụng tăng lên chính xác đúng số tiền vừa nạp.", "• Bước 5: ")

    add_heading(doc, "5.2. KOC rút tiền về tài khoản ngân hàng", 2)
    add_bullet(doc, "Vào menu 'Ví' (#/wallet) trên tài khoản KOC.", "• Bước 1: ")
    add_bullet(doc, "Xem số dư tiền kiếm được từ các đơn hàng ➔ Bấm nút 'Rút tiền'.", "• Bước 2: ")
    add_bullet(doc, "Nhập số tiền muốn rút (lưu ý số tiền tối thiểu là 10.000 VNĐ).", "• Bước 3: ")
    add_bullet(doc, "Bấm nút 'Nhận mã OTP' ➔ Kiểm tra hòm thư email để lấy mã xác thực 6 số ➔ Nhập vào ô mã OTP.", "• Bước 4: ")
    add_bullet(doc, "Bấm 'Xác nhận rút tiền'.", "• Bước 5: ")
    add_bullet(doc, "Kiểm tra kết quả: Hệ thống báo yêu cầu rút tiền đang được xử lý, số dư ví KOC bị trừ đúng số tiền yêu cầu, và lịch sử giao dịch hiển thị dòng 'Rút tiền - Đang xử lý'.", "• Bước 6: ")

    add_table(
        doc,
        ["Tình huống test kiểm tra bắt lỗi", "Thao tác của Tester", "Kết quả mong đợi"],
        [
            ("Rút số tiền dưới 10.000 VNĐ", "Nhập số tiền 5.000 VNĐ và bấm tiếp tục", "Hệ thống báo lỗi: Số tiền rút tối thiểu là 10.000 VNĐ; không cho gửi yêu cầu."),
            ("Rút số tiền lớn hơn số dư ví", "Ví có 100.000 VNĐ nhưng nhập rút 200.000 VNĐ", "Hệ thống báo lỗi: Số dư ví khả dụng không đủ."),
            ("Nhập sai mã OTP", "Nhập mã OTP bất kỳ không khớp với email", "Hệ thống báo lỗi mã OTP không chính xác; tiền trong ví giữ nguyên.")
        ],
        [2800, 3160, 3400],
        center_cols=[0]
    )

    # ----------------------------------------------------
    # PHẦN 6: LUỒNG 5 — TIẾP THỊ LIÊN KẾT (AFFILIATE)
    # ----------------------------------------------------
    add_heading(doc, "6. Luồng 5 — Tiếp thị liên kết (Affiliate & Hoa hồng)", 1)
    
    add_para(doc, "Mục tiêu: KOC nhận link sản phẩm để gắn vào bài đăng hoặc bio, hệ thống ghi nhận hoa hồng theo từng đơn bán thành công.")

    add_bullet(doc, "Doanh nghiệp gửi yêu cầu Affiliate: Khi tạo booking, Doanh nghiệp chọn hình thức 'Affiliate' hoặc 'Combo', điền link sản phẩm trên sàn TMĐT và tỷ lệ hoa hồng chia sẻ (ví dụ: 15%).", "• Bước 1: ")
    add_bullet(doc, "KOC nhận link bán hàng độc quyền: KOC đồng ý nhận đơn ➔ Vào menu 'Hoa hồng bán hàng' (#/affiliate) sẽ thấy Link tiếp thị (Tracking Link) riêng biệt được sinh tự động.", "• Bước 2: ")
    add_bullet(doc, "Kiểm tra theo dõi đơn hàng: Màn hình hiển thị danh sách đơn phát sinh, doanh số (GMV) và tiền hoa hồng KOC nhận được tương ứng với 3 trạng thái:", "• Bước 3: ")

    add_bullet(doc, "1. 'Chờ đối soát' (Màu vàng): Đơn hàng vừa phát sinh, hoa hồng ở mức tạm tính (chưa rút được).", "   - ")
    add_bullet(doc, "2. 'Đã xác nhận' (Màu xanh dương): Đơn hàng giao thành công và hết thời hạn đổi trả.", "   - ")
    add_bullet(doc, "3. 'Đã chi trả' (Màu xanh lá): Tiền hoa hồng được cộng chính thức vào Ví KOC để có thể rút về tài khoản ngân hàng.", "   - ")

    # ----------------------------------------------------
    # PHẦN 7: LUỒNG 6 — CHIẾN DỊCH LỚN (MEGA CAMPAIGN)
    # ----------------------------------------------------
    add_heading(doc, "7. Luồng 6 — Chiến dịch lớn (Mega Campaign nhiều KOC)", 1)
    
    add_para(doc, "Mục tiêu: Nhãn hàng muốn triển khai một chiến dịch lớn gồm hàng chục KOC cùng đăng bài trong một đợt ra mắt sản phẩm mới.")

    add_step(doc, "Bước 1", "Doanh nghiệp khởi tạo chiến dịch",
             "• Doanh nghiệp vào menu 'Chiến dịch lớn' (#/campaigns) ➔ Bấm nút 'Tạo chiến dịch mới'.\n"
             "• Điền thông tin: Tên chiến dịch, Mục tiêu, Ngành hàng, Tổng ngân sách dự kiến (ví dụ: 50.000.000 VNĐ), Số lượng KOC mong muốn tuyển (ví dụ: 10 bạn), Hạn chót hoàn thành ➔ Bấm 'Gửi yêu cầu'.")

    add_step(doc, "Bước 2", "Admin gửi Báo giá & Doanh nghiệp ký quỹ",
             "• Admin vào mục 'Điều phối chiến dịch' (#/campaigns), kiểm tra yêu cầu và nhập mức phí quản lý ➔ Bấm 'Gửi báo giá'.\n"
             "• Doanh nghiệp mở chiến dịch, xem báo giá minh bạch ➔ Bấm nút 'Đồng ý & Ký quỹ ngân sách'. Tiền ngân sách được đưa vào quỹ bảo đảm.")

    add_step(doc, "Bước 3", "Admin phân bổ KOC & Nghiệm thu chiến dịch",
             "• Admin chọn và thêm danh sách các bạn KOC phù hợp vào chiến dịch sao cho vừa vặn ngân sách.\n"
             "• KOC nhận lời mời và nộp nội dung bài đăng.\n"
             "• Doanh nghiệp duyệt bài và Admin bấm giải ngân cho từng bạn KOC theo tiến độ.")

    # ----------------------------------------------------
    # PHẦN 8: LUỒNG 7 — BOOKING AI CLONE AVATAR (NGƯỜI MẪU ẢO AI)
    # ----------------------------------------------------
    add_heading(doc, "8. Luồng 7 — Booking AI Clone Avatar (Người mẫu ảo AI)", 1)
    
    add_para(doc, "Mục tiêu: Đặt làm video quảng cáo bằng bản sao người mẫu ảo AI của KOC, giúp nhãn hàng tiết kiệm thời gian quay chụp.")

    add_bullet(doc, "KOC bật tính năng AI Clone: KOC vào menu 'AI Clone' (#/aiclone), đọc điều khoản về khai thác hình ảnh và bấm kích hoạt tham gia.", "• Bước 1: ")
    add_bullet(doc, "Doanh nghiệp đặt làm video AI: Doanh nghiệp vào menu 'Booking AI Clone' (#/aiclone-booking), chọn gương mặt KOC ảo, nhập kịch bản (Script) và yêu cầu bối cảnh ➔ Bấm gửi yêu cầu.", "• Bước 2: ")
    add_bullet(doc, "Admin báo giá & Sản xuất: Admin tiếp nhận, gửi báo giá chi phí sản xuất ➔ Doanh nghiệp duyệt chi phí ➔ Đội ngũ kỹ thuật dựng video AI và tải video lên hệ thống.", "• Bước 3: ")
    add_bullet(doc, "Quy trình duyệt 2 lớp an toàn: Doanh nghiệp vào duyệt video trước ➔ Sau đó KOC xem video để duyệt quyền sử dụng hình ảnh ➔ Đơn hàng hoàn tất và giải ngân.", "• Bước 4: ")

    # ----------------------------------------------------
    # PHẦN 9: LUỒNG 8 — DỊCH VỤ KOL / NGHỆ SĨ NỔI TIẾNG
    # ----------------------------------------------------
    add_heading(doc, "9. Luồng 8 — Dịch vụ KOL / Nghệ sĩ nổi tiếng (Dịch vụ VIP)", 1)
    
    add_para(doc, "Mục tiêu: Doanh nghiệp đặt lịch với các Celeb / Nghệ sĩ tên tuổi lớn với sự hỗ trợ kết nối trực tiếp từ sàn.")

    add_bullet(doc, "Doanh nghiệp gửi yêu cầu: Vào mục 'KOL / Nghệ sĩ' (#/kol), chọn nghệ sĩ mong muốn, nhập tóm tắt chiến dịch (Brief) và mức ngân sách dự kiến ➔ Bấm 'Gửi yêu cầu'.", "• Bước 1: ")
    add_bullet(doc, "Admin kết nối & Báo giá trọn gói: Admin liên hệ trực tiếp với quản lý của Nghệ sĩ, nhập báo giá chi tiết và thời gian làm việc gửi cho Doanh nghiệp.", "• Bước 2: ")
    add_bullet(doc, "Ký quỹ & Nghiệm thu: Doanh nghiệp duyệt báo giá và ký quỹ ➔ Nghệ sĩ quay chụp theo lịch ➔ Admin giao sản phẩm nghiệm thu ➔ Doanh nghiệp bấm 'Nghiệm thu hoàn tất'.", "• Bước 3: ")

    # ----------------------------------------------------
    # PHẦN 10: LUỒNG 9 — KHIẾU NẠI, TRANH CHẤP & THÔNG BÁO
    # ----------------------------------------------------
    add_heading(doc, "10. Luồng 9 — Khiếu nại, Tranh chấp & Thông báo hệ thống", 1)
    
    add_heading(doc, "10.1. Xử lý khiếu nại đơn hàng", 2)
    add_para(doc, "Khi xảy ra sự cố (ví dụ: KOC nộp video trễ hạn, hoặc Doanh nghiệp từ chối duyệt video vô lý):")
    add_bullet(doc, "Một trong hai bên bấm nút 'Gửi khiếu nại' ngay trên trang chi tiết đơn hàng, nhập rõ nguyên nhân và tải ảnh chụp bằng chứng.", "• Bước 1: ")
    add_bullet(doc, "Đơn hàng chuyển sang trạng thái 'Đang khiếu nại'. Tiền trong quỹ tạm giữ bị khóa lại để bảo vệ quyền lợi 2 bên.", "• Bước 2: ")
    add_bullet(doc, "Admin vào mục 'Khiếu nại' (#/complaints), xem ý kiến 2 bên và đưa ra phán quyết: Hoàn tiền lại cho Doanh nghiệp (nếu KOC vi phạm) hoặc Giải ngân cho KOC (nếu KOC đã làm đúng cam kết).", "• Bước 3: ")

    add_heading(doc, "10.2. Hệ thống thông báo (Quả chuông)", 2)
    add_para(doc, "Kiểm tra biểu tượng Quả chuông ở góc trên màn hình:")
    add_bullet(doc, "Khi có đơn hàng mới, video mới được nộp, có phản hồi sửa đổi hoặc tiền về ví, biểu tượng chuông phải hiện chấm đỏ kèm số lượng tin mới.", "• ")
    add_bullet(doc, "Bấm vào chuông: Xem danh sách thông báo và bấm vào từng tin sẽ chuyển hướng trực tiếp đến đúng trang đơn hàng liên quan.", "• ")

    # ----------------------------------------------------
    # PHẦN 11: BẢNG CHECKLIST KIỂM THỬ DÀNH CHO MARKETING
    # ----------------------------------------------------
    add_heading(doc, "11. Bảng Checklist kiểm thử nhanh (Dành riêng cho Marketing)", 1)
    
    add_para(doc, "Bạn có thể in bảng này ra hoặc đánh dấu trực tiếp trên file để theo dõi tiến độ kiểm thử của mình:")

    add_table(
        doc,
        ["STT", "Tính năng cần test", "Vai trò", "Thao tác chính", "Kết quả mong đợi", "Đánh giá"],
        [
            ("1", "Đăng ký KOC mới", "KOC", "Điền email, nhận OTP, điền thông tin cá nhân", "Xác thực OTP thành công, chuyển sang bước tạo hồ sơ", "[  ] Đạt"),
            ("2", "Khai báo kênh mạng xã hội", "KOC", "Nhập link kênh và tổng số follower", "Hệ thống ghi nhận thông tin để Admin kiểm duyệt", "[  ] Đạt"),
            ("3", "Admin duyệt KOC", "Admin", "Vào #/queue, xem chi tiết và bấm Duyệt", "KOC đăng nhập được và xuất hiện trên chợ #/find", "[  ] Đạt"),
            ("4", "Đăng ký Doanh nghiệp", "Brand", "Điền thông tin công ty, MST, tải giấy phép KD", "Gửi thành công, trạng thái chờ Admin duyệt", "[  ] Đạt"),
            ("5", "Admin duyệt Doanh nghiệp", "Admin", "Vào #/businesses, tìm tên công ty và bấm Duyệt", "Doanh nghiệp đăng nhập được vào #/dashboard", "[  ] Đạt"),
            ("6", "Tìm KOC & Tạo Booking", "Brand", "Vào #/find, chọn KOC, nhập brief, deadline", "Đơn tạo thành công; Tiền vào quỹ tạm giữ Escrow", "[  ] Đạt"),
            ("7", "KOC nhận việc", "KOC", "Vào #/bookings, xem đơn và bấm Nhận việc", "Trạng thái đơn chuyển sang Đang thực hiện", "[  ] Đạt"),
            ("8", "KOC nộp video nháp", "KOC", "Tải video demo/nháp lên hệ thống", "Trạng thái đơn chuyển sang Chờ Brand duyệt video", "[  ] Đạt"),
            ("9", "Brand yêu cầu sửa video", "Brand", "Bấm 'Yêu cầu sửa', nhập góp ý chi tiết", "KOC thấy góp ý và có nút tải video mới", "[  ] Đạt"),
            ("10", "Brand duyệt video", "Brand", "Bấm nút 'Duyệt video'", "Trạng thái chuyển sang Chờ KOC đăng bài", "[  ] Đạt"),
            ("11", "KOC nộp link bài đăng", "KOC", "Đăng bài lên TikTok/FB, dán link bài viết", "Trạng thái chuyển sang Đã đăng, chờ nghiệm thu", "[  ] Đạt"),
            ("12", "Brand hoàn tất & Đánh giá", "Brand", "Bấm 'Hoàn tất đơn hàng', chấm sao & nhận xét", "KOC nhận tiền vào ví; Đánh giá sao hiện trên hồ sơ", "[  ] Đạt"),
            ("13", "Nạp tiền ví Doanh nghiệp", "Brand", "Vào #/wallet, tạo mã nạp qua PayOS", "Quét QR xong, số dư khả dụng tăng lên ngay", "[  ] Đạt"),
            ("14", "KOC rút tiền về ngân hàng", "KOC", "Vào #/wallet, bấm Rút tiền, nhập OTP", "Số dư ví trừ đúng tiền, lệnh ở trạng thái Đang xử lý", "[  ] Đạt"),
            ("15", "Tiếp thị liên kết (Affiliate)", "KOC", "Nhận đơn affiliate, lấy tracking link", "Link tiếp thị hoạt động, theo dõi được hoa hồng", "[  ] Đạt")
        ],
        [500, 1600, 900, 2400, 2960, 1000],
        center_cols=[0, 2, 5]
    )

    # ----------------------------------------------------
    # PHẦN 12: HƯỚNG DẪN BÁO CÁO LỖI THÂN THIỆN
    # ----------------------------------------------------
    add_heading(doc, "12. Hướng dẫn gửi phản hồi & Báo cáo lỗi (Dành cho Marketing)", 1)
    
    add_para(doc, "Khi gặp bất kỳ chỗ nào bị kẹt, nút bấm không phản hồi, hoặc giao diện hiển thị khó hiểu, bạn không cần phải dùng từ ngữ kỹ thuật phức tạp. Chỉ cần gửi cho đội ngũ kỹ thuật một tin nhắn theo mẫu 5 dòng đơn giản sau:")

    add_callout(
        doc,
        "MẪU BÁO CÁO LỖI / GÓP Ý TRẢI NGHIỆM ĐƠN GIẢN",
        "1. Đường link trang bạn đang đứng: (Ví dụ: https://kocviet.com/#/find)\n"
        "2. Tài khoản bạn đang đăng nhập: (Ví dụ: Tài khoản Doanh nghiệp hay KOC? Email là gì?)\n"
        "3. Bạn vừa bấm nút gì hoặc làm thao tác gì? (Ví dụ: Vừa bấm nút 'Xác nhận Booking')\n"
        "4. Hiện tượng bạn thấy trên màn hình: (Ví dụ: Nút quay tròn hoài không dừng, hoặc màn hình hiện thông báo đỏ 'Lỗi kết nối')\n"
        "5. Ảnh chụp toàn màn hình: (Chụp rõ thông báo hoặc chỗ hiển thị lạ kèm theo tin nhắn).",
        fill=LIGHT_BLUE,
        border_color=BLUE
    )

    add_para(doc, "💡 Mẹo chụp ảnh màn hình nhanh trên Windows: Nhấn tổ hợp phím Windows + Shift + S, quét chọn vùng màn hình bị lỗi, sau đó mở khung chat (Zalo/Telegram/Slack) và bấm Ctrl + V để dán ảnh gửi ngay cho đội kỹ thuật.", bold_prefix="💡 Mẹo chụp ảnh màn hình nhanh trên Windows: ", italic=True)

    # Save document
    doc.save(str(OUTPUT_DOCX))
    print(f"Marketing test guide successfully written to: {OUTPUT_DOCX}")

if __name__ == "__main__":
    build_marketing_guide()
