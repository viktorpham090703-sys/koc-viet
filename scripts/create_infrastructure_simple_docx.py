from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from create_infrastructure_docx import (
    BLACK,
    BLUE,
    LIGHT_BLUE,
    MID_GRAY,
    NAVY,
    PALE_BLUE,
    TABLE_WIDTH,
    add_page_number,
    set_cell_margins,
    set_cell_shading,
    set_repeat_table_header,
    set_run_font,
    set_table_geometry,
)


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "infrastructure-plan-simple.docx"


def configure(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, NAVY, 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167

    header = section.header.paragraphs[0]
    set_run_font(header.add_run("KOC VIỆT  |  HẠ TẦNG DỄ HIỂU"), size=8.5, color=MID_GRAY, bold=True)
    add_page_number(section.footer.paragraphs[0])


def add_title_block(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(28)
    p.paragraph_format.space_after = Pt(5)
    set_run_font(p.add_run("KẾ HOẠCH HẠ TẦNG KOC VIỆT"), size=27, color=NAVY, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    set_run_font(p.add_run("Bản giải thích dành cho người không chuyên kỹ thuật"), size=14, color=BLUE)

    add_callout(
        doc,
        "Kết luận ngắn",
        "Hiện tại có thể dùng AWS. Nếu Bộ yêu cầu máy chủ và dữ liệu phải nằm tại Việt Nam, chuyển sang Viettel IDC. Hai phương án làm cùng một công việc; khác nhau chủ yếu ở nơi đặt máy chủ, cách tính phí và thủ tục quản lý.",
        fill="EAF3FB",
    )


def add_callout(doc, label, text, fill="FFF4E5"):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.15)
    p.paragraph_format.right_indent = Inches(0.15)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(10)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    set_run_font(p.add_run(f"{label}: "), size=11, color=NAVY, bold=True)
    set_run_font(p.add_run(text), size=11, color=BLACK)


def add_heading(doc, text, level=1):
    doc.add_paragraph(text, style=f"Heading {level}")


def add_para(doc, text, bold_start=None):
    p = doc.add_paragraph()
    if bold_start and text.startswith(bold_start):
        set_run_font(p.add_run(bold_start), size=11, color=BLACK, bold=True)
        set_run_font(p.add_run(text[len(bold_start):]), size=11, color=BLACK)
    else:
        set_run_font(p.add_run(text), size=11, color=BLACK)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        set_run_font(p.add_run(item), size=11, color=BLACK)


def add_steps(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        set_run_font(p.add_run(item), size=11, color=BLACK)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        cell = table.cell(0, idx)
        cell.text = ""
        set_cell_shading(cell, BLUE)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        set_run_font(p.add_run(header), size=10, color="FFFFFF", bold=True)
    set_repeat_table_header(table.rows[0])

    for r_idx, row in enumerate(rows, start=1):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = ""
            if r_idx % 2 == 0:
                set_cell_shading(cell, LIGHT_BLUE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            set_run_font(p.add_run(value), size=9.7, color=BLACK)
            set_cell_margins(cell, top=100, bottom=100, start=120, end=120)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def build():
    doc = Document()
    configure(doc)
    add_title_block(doc)

    add_heading(doc, "1. Hệ thống này cần những gì?", 1)
    add_para(doc, "Để website hoạt động, chúng ta cần ba nhóm dịch vụ chính. Có thể hiểu chúng giống như một văn phòng trực tuyến:")
    add_table(
        doc,
        ["Thành phần", "Hiểu đơn giản", "Dùng để làm gì"],
        [
            ("Máy chủ", "Một máy tính được thuê trên Internet", "Chạy website và phần xử lý phía sau"),
            ("Cơ sở dữ liệu", "Tủ hồ sơ điện tử", "Lưu tài khoản, booking, số dư, escrow và lịch sử giao dịch"),
            ("Kho ảnh", "Kho chứa ảnh trực tuyến", "Lưu avatar, logo và ảnh bìa của KOC/doanh nghiệp"),
        ],
        [1900, 3100, 4360],
    )
    add_callout(doc, "Lưu ý về video", "Hệ thống không lưu file video. Hệ thống chỉ lưu đường link dẫn đến video trên Facebook, TikTok, YouTube hoặc nền tảng khác.")

    add_heading(doc, "2. Phương án đang dùng: AWS", 1)
    add_para(doc, "AWS là nền tảng điện toán đám mây quốc tế. Máy chủ dự kiến đặt tại Singapore, gần Việt Nam.")
    add_table(
        doc,
        ["Tên AWS", "Vai trò", "Cấu hình dự kiến"],
        [
            ("EC2", "Máy tính chạy website", "2 CPU, 4 GB RAM"),
            ("RDS PostgreSQL", "Tủ hồ sơ điện tử", "Database tách riêng, có backup"),
            ("S3", "Kho ảnh", "Chỉ lưu ảnh, không lưu video"),
        ],
        [2200, 3300, 3860],
    )
    add_para(doc, "Ứng dụng được đóng gói bằng Docker. Có thể hiểu Docker là một chiếc hộp chứa sẵn ứng dụng và các phần cần thiết, giúp cài đặt lại dễ dàng hơn.")

    add_heading(doc, "3. AWS dự kiến tốn bao nhiêu mỗi tháng?", 1)
    add_para(doc, "Các số dưới đây chỉ là ước tính, không phải báo giá cố định. Hóa đơn thật phụ thuộc vào lượng sử dụng và giá AWS tại thời điểm thanh toán.")
    add_table(
        doc,
        ["Khoản phí", "Ước tính mỗi tháng", "Giải thích"],
        [
            ("Máy chủ EC2", "$30-$40", "Chạy website liên tục 24/7"),
            ("Ổ đĩa và IP", "$6-$8", "Ổ cứng và địa chỉ IP cố định"),
            ("RDS PostgreSQL", "$21-$36", "Database, ổ đĩa và backup cơ bản"),
            ("S3 lưu ảnh", "$0.25-$1", "Giai đoạn đầu có ít ảnh"),
            ("Tổng dự kiến", "$58-$85", "Nên dự phòng ngân sách $65-$95"),
        ],
        [2600, 2200, 4560],
    )
    add_callout(doc, "Điều cần nhớ", "AWS tính phí bằng USD. Chi phí có thể tăng khi có nhiều người dùng, nhiều ảnh, nhiều log hoặc khi phải nâng máy chủ.", fill="FFF4E5")

    add_heading(doc, "4. Khi nào phải dùng máy chủ trong nước?", 1)
    add_para(doc, "Không phải cứ làm dự án cho Bộ là tự động bắt buộc dùng máy chủ trong nước. Tuy nhiên, phải chuyển sang hạ tầng trong nước nếu hợp đồng hoặc đơn vị phụ trách CNTT yêu cầu một trong các nội dung sau:")
    add_bullets(doc, [
        "Máy chủ và dữ liệu phải đặt tại Việt Nam.",
        "Không được chuyển dữ liệu cá nhân hoặc dữ liệu nghiệp vụ ra nước ngoài.",
        "Phải sử dụng nhà cung cấp nằm trong danh sách được cơ quan chủ quản chấp thuận.",
        "Database, ảnh và bản backup đều phải nằm tại trung tâm dữ liệu trong nước.",
    ])
    add_callout(doc, "Khuyến nghị", "Trước khi mua dịch vụ dài hạn, cần xin xác nhận bằng email hoặc văn bản từ đơn vị phụ trách CNTT của Bộ.", fill="FDECEC")

    add_heading(doc, "5. Phương án trong nước: Viettel IDC", 1)
    add_para(doc, "Viettel IDC cung cấp các dịch vụ tương đương AWS nhưng trung tâm dữ liệu đặt tại Việt Nam.")
    add_table(
        doc,
        ["Đang dùng trên AWS", "Chuyển sang Viettel IDC", "Công việc không thay đổi"],
        [
            ("EC2", "Viettel Cloud Server", "Chạy website và backend"),
            ("RDS PostgreSQL", "Viettel Database Service", "Lưu dữ liệu hệ thống"),
            ("S3", "Viettel Cloud Object Storage", "Lưu ảnh"),
        ],
        [2500, 3400, 3460],
    )
    add_para(doc, "Cấu hình nên yêu cầu Viettel IDC:")
    add_bullets(doc, [
        "Cloud Server: 2 CPU, 4 GB RAM và ổ đĩa 30-50 GB.",
        "PostgreSQL 15: database tách riêng, chỉ kết nối nội bộ với máy chủ.",
        "Backup database mỗi ngày và giữ ít nhất 7 ngày.",
        "Object Storage: lưu dưới 10 GB ảnh trong giai đoạn đầu.",
        "Một địa chỉ IP cố định, firewall và theo dõi hoạt động hệ thống.",
    ])

    add_heading(doc, "6. Viettel IDC có rẻ hơn AWS không?", 1)
    add_para(doc, "Chưa thể khẳng định. Viettel IDC chưa công khai đầy đủ giá cho toàn bộ cấu hình trên; cần xin báo giá chính thức.")
    add_table(
        doc,
        ["AWS", "Viettel IDC"],
        [
            ("Ước tính $58-$85 mỗi tháng", "Phải xin báo giá"),
            ("Máy chủ đặt tại Singapore", "Máy chủ và dữ liệu có thể đặt tại Việt Nam"),
            ("Thanh toán quốc tế bằng USD", "Thuận tiện hơn cho hợp đồng và hỗ trợ trong nước"),
            ("Tài liệu kỹ thuật rất đầy đủ", "Có hỗ trợ tiếng Việt và đầu mối trong nước"),
        ],
        [4680, 4680],
    )
    add_para(doc, "Khi so giá, phải yêu cầu Viettel ghi rõ máy chủ, ổ đĩa, IP, băng thông, database, backup, kho ảnh, hỗ trợ, SLA và thuế. Không nên chỉ so giá máy chủ 4 GB RAM.")

    add_heading(doc, "7. GitHub Actions và Docker có tiếp tục dùng được không?", 1)
    add_para(doc, "Có. Việc chuyển từ AWS sang Viettel IDC không bắt buộc phải viết lại toàn bộ ứng dụng.")
    add_bullets(doc, [
        "GitHub Actions tiếp tục kiểm tra và đóng gói ứng dụng.",
        "Docker tiếp tục dùng để chạy frontend, backend và Nginx.",
        "Phần database chỉ đổi địa chỉ kết nối.",
        "Phần lưu ảnh chỉ đổi địa chỉ kho ảnh và khóa truy cập.",
    ])
    add_callout(doc, "Ngoại lệ", "Nếu Bộ không cho phép kết nối trực tiếp với GitHub hoặc kho Docker nước ngoài, có thể dùng máy triển khai nội bộ và Viettel Container Registry.")

    add_heading(doc, "8. Cách ra quyết định", 1)
    add_steps(doc, [
        "Hỏi Bộ có bắt buộc đặt máy chủ và dữ liệu tại Việt Nam hay không.",
        "Nếu không bắt buộc, có thể tiếp tục dùng AWS theo cấu hình hiện tại.",
        "Nếu bắt buộc, xin báo giá Viettel IDC cho đúng cấu hình trong tài liệu này.",
        "Triển khai môi trường thử nghiệm trước, không dùng dữ liệu thật.",
        "Kiểm tra đăng nhập, OTP, upload avatar, booking, ví và escrow.",
        "Chỉ chuyển hệ thống chính thức sau khi backup và có phương án quay lại nếu xảy ra lỗi.",
    ])

    add_heading(doc, "9. Những điều cần xác nhận trước khi ký", 1)
    add_bullets(doc, [
        "Máy chủ, database, ảnh và backup nằm ở quốc gia nào?",
        "Chi phí hàng tháng đã gồm IP, ổ đĩa, băng thông và backup chưa?",
        "Database có được backup tự động và có thể khôi phục không?",
        "Thời gian cam kết hệ thống hoạt động (SLA) là bao nhiêu?",
        "Ai hỗ trợ khi website hoặc database gặp sự cố?",
        "Có cho phép dùng GitHub Actions và kho Docker bên ngoài không?",
        "Ai chịu trách nhiệm phê duyệt hạ tầng và an toàn thông tin?",
    ])

    add_heading(doc, "10. Kết luận", 1)
    add_callout(
        doc,
        "Phương án đề xuất",
        "Tiếp tục chuẩn bị hệ thống theo AWS khi chưa có yêu cầu khác. Nếu Bộ yêu cầu toàn bộ dữ liệu đặt tại Việt Nam, sử dụng Viettel Cloud Server, Viettel Database Service PostgreSQL và Viettel Cloud Object Storage. Không ký hợp đồng dài hạn với nhà cung cấp trước khi có xác nhận về nơi đặt dữ liệu.",
        fill="EAF3FB",
    )

    add_heading(doc, "Nguồn tham khảo", 1)
    for label, url in [
        ("AWS Pricing Calculator", "https://calculator.aws/"),
        ("Viettel Cloud Server", "https://viettelidc.com.vn/en/cloud-server"),
        ("Viettel IDC - yêu cầu báo giá", "https://www.viettelidc.com.vn/en/Home/PriceList"),
    ]:
        p = doc.add_paragraph(style="List Bullet")
        set_run_font(p.add_run(f"{label}: "), size=10, color=BLACK, bold=True)
        set_run_font(p.add_run(url), size=10, color=BLUE)

    doc.core_properties.title = "Kế hoạch hạ tầng KOC Việt - Bản dễ hiểu"
    doc.core_properties.subject = "Giải thích AWS và Viettel IDC dành cho người không chuyên kỹ thuật"
    doc.core_properties.author = "KOC Việt"
    doc.core_properties.keywords = "AWS, Viettel IDC, hạ tầng, bản dễ hiểu"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
